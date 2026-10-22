"""
Khung viền trang và Trang bìa (Bìa ngoài & Bìa phụ) chuẩn Đại học Công nghệ Thông tin (UIT).
Tuân thủ cấu trúc chuẩn học thuật, hai trang đầu đồng bộ thứ tự nội dung, không màu nền, không viền trang dư thừa.
Tuân thủ Invariant 01, Invariant 02 (< 400 dòng) và Invariant 16.
"""
import os
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from uit_report_content import UIT_FRONT_MATTER, format_table_cell, set_table_borders


def clear_page_border(section):
    """Hủy bỏ hoàn toàn khung viền trang cho section theo đúng chuẩn Word học thuật UIT."""
    sectPr = section._sectPr
    for pb in sectPr.xpath('.//w:pgBorders'):
        sectPr.remove(pb)


def _add_cover_header(doc):
    """Thêm tiêu đề cơ quan chủ quản và đơn vị đào tạo chuẩn UIT."""
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.line_spacing = 1.2
    p_header.paragraph_format.space_before = Pt(0)
    p_header.paragraph_format.space_after = Pt(2)

    for txt, sz, bld in [
        (UIT_FRONT_MATTER["ministry"] + "\n", 12, True),
        (UIT_FRONT_MATTER["vnuhcm"] + "\n", 12, True),
        (UIT_FRONT_MATTER["univ"] + "\n", 13, True),
        (UIT_FRONT_MATTER["faculty"] + "\n", 12, True)
    ]:
        r = p_header.add_run(txt)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(sz)
        r.font.bold = bld
        r.font.color.rgb = RGBColor(0, 0, 0)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(6)
    r_line = p_line.add_run("---------------------------------------------")
    r_line.font.name = 'Times New Roman'
    r_line.font.bold = True
    r_line.font.color.rgb = RGBColor(0, 0, 0)


def _add_cover_logo(doc, assets_dir):
    """Thêm logo chính thức của UIT."""
    logo_path = os.path.join(assets_dir, "uit_logo.png")
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(4)
        p_logo.paragraph_format.space_after = Pt(8)
        p_logo.add_run().add_picture(logo_path, width=Inches(1.35))


def _add_cover_title_block(doc):
    """Thêm khối loại báo cáo, tên đề tài tiếng Việt và tiếng Anh."""
    p_rep_type = doc.add_paragraph()
    p_rep_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rep_type.paragraph_format.space_before = Pt(4)
    p_rep_type.paragraph_format.space_after = Pt(3)
    r_rt = p_rep_type.add_run(UIT_FRONT_MATTER["report_type"] + "\n" + UIT_FRONT_MATTER["course_name"])
    r_rt.font.name = 'Times New Roman'
    r_rt.font.size = Pt(12)
    r_rt.font.bold = True
    r_rt.font.color.rgb = RGBColor(0, 0, 0)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 1.25
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(2)

    r_pre = p_title.add_run("ĐỀ TÀI:\n")
    r_pre.font.name = 'Times New Roman'
    r_pre.font.size = Pt(13)
    r_pre.font.bold = True
    r_pre.font.color.rgb = RGBColor(0, 0, 0)

    r_title = p_title.add_run(UIT_FRONT_MATTER["title_vi"] + "\n")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 0, 0)

    r_en = p_title.add_run(UIT_FRONT_MATTER["title_en"] + "\n")
    r_en.font.name = 'Times New Roman'
    r_en.font.size = Pt(10.5)
    r_en.font.italic = True
    r_en.font.color.rgb = RGBColor(0, 0, 0)

    r_maj = p_title.add_run(UIT_FRONT_MATTER["major"])
    r_maj.font.name = 'Times New Roman'
    r_maj.font.size = Pt(12)
    r_maj.font.bold = True
    r_maj.font.color.rgb = RGBColor(0, 0, 0)


def add_outer_cover_page(doc, assets_dir):
    """Tạo Trang Bìa Ngoài (Outer Cover) chuẩn UIT: Đồng bộ bố cục với bìa trong."""
    _add_cover_header(doc)
    _add_cover_logo(doc, assets_dir)
    _add_cover_title_block(doc)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(14)

    # Khối thông tin Giảng viên hướng dẫn & Sinh viên thực hiện
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_info.paragraph_format.line_spacing = 1.3
    p_info.paragraph_format.space_before = Pt(4)
    p_info.paragraph_format.space_after = Pt(4)

    runs = [
        ("GIẢNG VIÊN HƯỚNG DẪN: ", True),
        (f"{UIT_FRONT_MATTER['instructor_name']}\n", True),
        ("CÁN BỘ PHẢN BIỆN: ", True),
        (f"{UIT_FRONT_MATTER['reviewer_name']}\n", False),
        ("SINH VIÊN THỰC HIỆN: ", True),
        (f"{UIT_FRONT_MATTER['student_name']} - MSSV: {UIT_FRONT_MATTER['student_id']}\n", True),
        ("LỚP: ", True),
        (f"{UIT_FRONT_MATTER['class_name']} - KHÓA HỌC: {UIT_FRONT_MATTER['cohort']}", False)
    ]
    for txt, bld in runs:
        r = p_info.add_run(txt)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.bold = bld
        r.font.color.rgb = RGBColor(0, 0, 0)

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(18)

    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_loc = p_loc.add_run(UIT_FRONT_MATTER["location"])
    r_loc.font.name = 'Times New Roman'
    r_loc.font.size = Pt(12)
    r_loc.font.bold = True
    r_loc.font.color.rgb = RGBColor(0, 0, 0)


def add_inner_cover_page(doc, assets_dir):
    """Tạo Trang Bìa Phụ (Inner Cover) chuẩn UIT: Có Logo UIT và Bảng thông tin GVHD/CBPB/Sinh viên."""
    _add_cover_header(doc)
    _add_cover_logo(doc, assets_dir)
    _add_cover_title_block(doc)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(8)

    col_widths = [Inches(2.2), Inches(4.1)]
    t_info = doc.add_table(rows=4, cols=2)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_info, col_widths=col_widths)

    info_rows = [
        (UIT_FRONT_MATTER["instructor_title"], UIT_FRONT_MATTER["instructor_name"], True),
        (UIT_FRONT_MATTER["reviewer_title"], UIT_FRONT_MATTER["reviewer_name"], False),
        (UIT_FRONT_MATTER["student_header"], f"{UIT_FRONT_MATTER['student_name']} - MSSV: {UIT_FRONT_MATTER['student_id']}", True),
        ("LỚP - KHÓA HỌC:", f"{UIT_FRONT_MATTER['class_name']} - {UIT_FRONT_MATTER['cohort']}", False)
    ]
    for i, (label, val, val_bold) in enumerate(info_rows):
        format_table_cell(t_info.rows[i].cells[0], label, bold=True, font_size=11)
        format_table_cell(t_info.rows[i].cells[1], val, bold=val_bold, font_size=11)

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(12)

    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_loc = p_loc.add_run(UIT_FRONT_MATTER["location"])
    r_loc.font.name = 'Times New Roman'
    r_loc.font.size = Pt(12)
    r_loc.font.bold = True
    r_loc.font.color.rgb = RGBColor(0, 0, 0)


def add_cover_page(doc, assets_dir, is_sub_cover=False):
    """Bộ điều phối trang bìa chính hoặc bìa phụ chuẩn UIT."""
    if is_sub_cover:
        add_inner_cover_page(doc, assets_dir)
    else:
        add_outer_cover_page(doc, assets_dir)
