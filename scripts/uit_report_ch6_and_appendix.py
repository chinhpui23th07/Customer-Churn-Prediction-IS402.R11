"""
Nội dung Chương 6, Tài liệu tham khảo IEEE và Phụ lục chuẩn UIT.
Tuân thủ Invariant 01, Invariant 02 (< 400 dòng) và Invariant 16.
"""
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from uit_report_content import (
    add_uit_body_p, add_uit_heading_1, add_uit_heading_2,
    IEEE_REFERENCES_LIST
)


def build_chapter_6(doc):
    """Chương 6: Kết luận và hướng phát triển."""
    add_uit_heading_1(doc, "CHƯƠNG 6: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    h1 = add_uit_heading_2(doc, "6.1. Các Kết Quả Đạt Được Của Đồ Án")
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(2)
    p1 = add_uit_body_p(doc, (
        "Sau quá trình nghiên cứu, thiết kế và triển khai thực nghiệm nghiêm túc, đề tài đã hoàn thành xuất sắc toàn bộ các mục tiêu: "
        "1) Xây dựng thành công nền tảng MLOps Level 2 hoàn chỉnh giải quyết triệt để bài toán Training-Serving Skew và Data Drift; "
        "2) Tiền xử lý và khử trùng lặp thành công bộ dữ liệu lớn 442,211 khách hàng, bảo đảm 100% Data Contracts với Great Expectations; "
        "3) Tinh chỉnh siêu tham số Bayes bằng Optuna đưa mô hình XGBoost đạt hiệu năng vượt trội ROC-AUC 0.9990; "
        "4) Cung cấp REST API hiệu năng cao trên nền FastAPI hỗ trợ dự đoán thời gian thực, dự đoán theo lô và A/B Testing deterministic; "
        "5) Trực quan hóa giải thích nguyên nhân rời bỏ bằng SHAP TreeExplainer gắn liền khuyến nghị kinh doanh."
    ))
    if p1:
        p1.paragraph_format.space_after = Pt(4)

    h2 = add_uit_heading_2(doc, "6.2. Phân Tích Giới Hạn Dữ Liệu Tổng Hợp Kaggle")
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(2)
    p2 = add_uit_body_p(doc, (
        "Báo cáo chỉ ra một phát hiện khoa học quan trọng: Tập dữ liệu Kaggle có sự thiên lệch mang tính quy luật tất định "
        "(khách hàng sử dụng hợp đồng theo tháng - Monthly Contract có tỷ lệ rời bỏ lên tới 99.84%). Mặc dù điều này giúp các mô hình "
        "học máy đạt độ chính xác gần như tuyệt đối (ROC-AUC ~ 0.999), trong môi trường sản xuất thực tế doanh nghiệp cần bổ sung các đặc trưng "
        "hành vi giao dịch vi mô và lịch sử thanh toán để tăng cường tính khái quát hóa."
    ))
    if p2:
        p2.paragraph_format.space_after = Pt(4)

    h3 = add_uit_heading_2(doc, "6.3. Hướng Phát Triển Tương Lai")
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(2)
    p3 = add_uit_body_p(doc, (
        "Trong các giai đoạn phát triển tiếp theo, đề tài có thể mở rộng theo các hướng: "
        "1) Nâng cấp Online Feature Store sang cụm phân tán Redis Cluster / Azure Cosmos DB để đáp ứng hàng trăm nghìn truy vấn đồng thời; "
        "2) Tích hợp học sâu dạng bảng (TabNet, FT-Transformer) và mô hình chuỗi thời gian phân tích hành vi theo tuần tự; "
        "3) Triển khai điều phối luồng công việc phức tạp bằng Apache Airflow hoặc Kubeflow Pipelines trên cụm Kubernetes (AKS/EKS); "
        "4) Tích hợp cơ chế bảo vệ quyền riêng tư người dùng theo chuẩn GDPR bằng kỹ thuật Differential Privacy."
    ))
    if p3:
        p3.paragraph_format.space_after = Pt(4)


def build_references_and_appendices(doc):
    """Tài liệu tham khảo chuẩn IEEE và Phụ lục A & B."""
    add_uit_heading_1(doc, "TÀI LIỆU THAM KHẢO CHUẨN IEEE")
    for ref in IEEE_REFERENCES_LIST:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_ref.paragraph_format.line_spacing = 1.25
        p_ref.paragraph_format.space_after = Pt(4)
        p_ref.paragraph_format.left_indent = Inches(0.39)
        p_ref.paragraph_format.first_line_indent = Inches(-0.39)
        r = p_ref.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
    doc.add_page_break()

    add_uit_heading_1(doc, "PHỤ LỤC")
    add_uit_heading_2(doc, "Phụ lục A: Hướng Dẫn Cài Đặt và Khởi Chạy Nhanh Hệ Thống")

    def _add_list_item(doc, text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)

    _add_list_item(doc, "1. Clone kho mã nguồn dự án: git clone <repo_url> && cd CLOUD")
    _add_list_item(doc, "2. Cài đặt các gói thư viện phụ thuộc: pip install -r requirements.txt")
    _add_list_item(doc, "3. Thực thi pipeline tiền xử lý và nạp Feature Store: python src/data/preprocess.py && python src/features/feature_store.py")
    _add_list_item(doc, "4. Chạy tối ưu hóa siêu tham số và huấn luyện mô hình: python src/models/train_pipeline.py")
    _add_list_item(doc, "5. Khởi chạy REST API và Web Serving: uvicorn src.api.main:app --host 0.0.0.0 --port 8000")
    _add_list_item(doc, "6. Chạy bộ kiểm thử tự động toàn diện: pytest tests/ -v")

    add_uit_heading_2(doc, "Phụ lục B: Danh Mục Các Endpoint REST API Chính Thức")
    _add_list_item(doc, "• GET /health : Kiểm tra trạng thái sẵn sàng của Model, Feast Online Store và MLflow Registry.")
    _add_list_item(doc, "• POST /predict : Dự đoán nguy cơ rời bỏ cho khách hàng đơn lẻ kèm giải thích SHAP XAI.")
    _add_list_item(doc, "• POST /predict/ab_test : Điều hướng tất định A/B Testing giữa Champion và Challenger.")
    _add_list_item(doc, "• POST /batch_predict : Dự đoán hàng loạt vector hóa hiệu năng cao kết hợp single-transaction DB log.")
    _add_list_item(doc, "• POST /upload_csv_predict : Dự đoán hàng loạt trực tiếp từ tệp CSV tải lên.")
    _add_list_item(doc, "• GET /features/{customer_id} : Truy vấn vector đặc trưng thời gian thực từ Feast Online Store.")
    _add_list_item(doc, "• GET /drift/status : Truy vấn báo cáo giám sát trôi dạt dữ liệu thời gian thực.")
    _add_list_item(doc, "• POST /api/v1/trigger_retrain : Webhook kích hoạt tái huấn luyện mô hình tự động.")
    _add_list_item(doc, "• GET /metrics : Endpoint cung cấp thông số vận hành hệ thống theo chuẩn Observability.")
