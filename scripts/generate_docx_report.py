"""
Trình tạo Báo Cáo Khóa Luận / Đồ Án chuẩn Đại học Công nghệ Thông tin (UIT - ĐHQG-HCM).
Điều phối khung viền bìa đôi, logo UIT, bìa phụ, biểu mẫu nhận xét GVHD/GVPB, hội đồng chấm,
lời cam đoan, lời cảm ơn, tóm tắt song ngữ, mục lục phân cấp, danh mục hình ảnh, bảng biểu,
6 chương chuyên sâu gắn số liệu thực nghiệm, tài liệu tham khảo IEEE và phụ lục.
"""
import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from uit_report_content import load_json_data
from uit_report_frontmatter import (
    clear_page_border,
    add_cover_page, add_review_forms,
    add_pledge_and_thanks, add_abstract_bilingual, add_toc_and_lists
)
from uit_report_ch1_to_3 import (
    build_chapter_1, build_chapter_2, build_chapter_3
)
from uit_report_ch4_to_6 import (
    build_chapter_4, build_chapter_5, build_chapter_6,
    build_references_and_appendices
)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(BASE_DIR, "docs")
ASSETS_DIR = os.path.join(BASE_DIR, "assets")
IMAGES_DIR = os.path.join(DOCS_DIR, "images")
DATA_PROCESSED_DIR = os.path.join(BASE_DIR, "data", "processed")
MODELS_DIR = os.path.join(BASE_DIR, "data", "models")
OUTPUT_DOCX = os.path.join(DOCS_DIR, "Bao_Cao_Bai_Tap_Lon_MLOps_Customer_Churn.docx")


def add_page_number_to_run(run):
    """Gắn trường số trang tự động vào run của footer."""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)


def build_uit_report():
    """Hàm chính điều phối lắp ráp toàn bộ tài liệu báo cáo khóa luận UIT."""
    os.makedirs(DOCS_DIR, exist_ok=True)
    doc = Document()

    # SECTION 1: Trang bìa duy nhất chuẩn UIT (Không header/footer, không viền thừa)
    sec1 = doc.sections[0]
    sec1.top_margin = Inches(0.984)
    sec1.bottom_margin = Inches(0.984)
    sec1.left_margin = Inches(1.181)
    sec1.right_margin = Inches(0.787)
    clear_page_border(sec1)

    add_cover_page(doc, ASSETS_DIR)

    # SECTION 2: Phần đầu Front Matter (Không viền, đánh số trang footer)
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.header.is_linked_to_previous = False
    sec2.footer.is_linked_to_previous = False
    sec2.top_margin = Inches(0.984)
    sec2.bottom_margin = Inches(0.984)
    sec2.left_margin = Inches(1.181)
    sec2.right_margin = Inches(0.787)
    clear_page_border(sec2)

    p_f2 = sec2.footer.paragraphs[0]
    p_f2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f2 = p_f2.add_run()
    r_f2.font.name = 'Times New Roman'
    r_f2.font.size = Pt(10)
    add_page_number_to_run(r_f2)

    sectPr2 = sec2._sectPr
    pgNumType2 = OxmlElement('w:pgNumType')
    pgNumType2.set(qn('w:start'), '1')
    pgNumType2.set(qn('w:fmt'), 'lowerRoman')
    sectPr2.append(pgNumType2)

    add_review_forms(doc)
    add_pledge_and_thanks(doc)
    add_abstract_bilingual(doc)
    add_toc_and_lists(doc)

    # Đọc dữ liệu artifacts thực nghiệm
    eval_data = load_json_data(os.path.join(MODELS_DIR, "evaluation_results.json"))
    io_data = load_json_data(os.path.join(DATA_PROCESSED_DIR, "io_benchmark_results.json"))
    lat_data = load_json_data(os.path.join(DATA_PROCESSED_DIR, "region_latency_results.json"))
    bias_data = load_json_data(os.path.join(DATA_PROCESSED_DIR, "dataset_bias_analysis.json"))
    tune_data = load_json_data(os.path.join(MODELS_DIR, "tuning_history.json"))
    err_data = load_json_data(os.path.join(DATA_PROCESSED_DIR, "error_analysis_report.json"))
    dq_data = load_json_data(os.path.join(DATA_PROCESSED_DIR, "data_quality_report.json"))

    # SECTION 3: Thân bài 6 Chương (Running Header, Footer số Ả Rập bắt đầu từ 1, không viền)
    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec3.header.is_linked_to_previous = False
    sec3.footer.is_linked_to_previous = False
    sec3.top_margin = Inches(0.984)
    sec3.bottom_margin = Inches(0.984)
    sec3.left_margin = Inches(1.181)
    sec3.right_margin = Inches(0.787)
    clear_page_border(sec3)

    # Đặt số trang bắt đầu từ 1
    sectPr3 = sec3._sectPr
    pgNumType3 = OxmlElement('w:pgNumType')
    pgNumType3.set(qn('w:start'), '1')
    pgNumType3.set(qn('w:fmt'), 'decimal')
    sectPr3.append(pgNumType3)

    # Running Header
    p_h3 = sec3.header.paragraphs[0]
    p_h3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_h3 = p_h3.add_run("Trường Đại học Công nghệ Thông tin - ĐHQG-HCM | Báo Cáo Khóa Luận Tốt Nghiệp")
    r_h3.font.name = 'Times New Roman'
    r_h3.font.size = Pt(9)
    r_h3.font.italic = True
    r_h3.font.color.rgb = RGBColor(0, 0, 0)

    # Running Footer
    p_f3 = sec3.footer.paragraphs[0]
    p_f3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f3 = p_f3.add_run()
    r_f3.font.name = 'Times New Roman'
    r_f3.font.size = Pt(11)
    add_page_number_to_run(r_f3)

    # Lắp ráp 6 Chương chuẩn UIT
    build_chapter_1(doc, IMAGES_DIR)
    doc.add_page_break()

    build_chapter_2(doc, IMAGES_DIR)
    doc.add_page_break()

    build_chapter_3(doc, IMAGES_DIR)
    doc.add_page_break()

    build_chapter_4(doc, IMAGES_DIR, eval_data, tune_data, dq_data, err_data)
    doc.add_page_break()

    build_chapter_5(doc, IMAGES_DIR, io_data, lat_data, bias_data)
    doc.add_page_break()

    build_chapter_6(doc)
    doc.add_page_break()

    build_references_and_appendices(doc)

    doc.save(OUTPUT_DOCX)
    print(f"[+] Báo cáo Khóa luận / Đồ án chuẩn UIT đã được tạo thành công tại: {OUTPUT_DOCX}")


if __name__ == "__main__":
    build_uit_report()
