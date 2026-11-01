"""
Phần đầu (Front Matter) của Báo Cáo Khóa Luận / Đồ Án chuẩn UIT:
- Phiếu nhận xét của Cán bộ hướng dẫn & Phản biện
- Thông tin Hội đồng chấm bảo vệ
- Lời cam đoan, Lời cảm ơn, Tóm tắt song ngữ
- Mục lục phân cấp, Danh mục hình ảnh, Danh mục bảng biểu và Danh mục từ viết tắt.
"""
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT

from uit_report_content import (
    UIT_FRONT_MATTER, ABBREVIATIONS_LIST, FIGURES_LIST, TABLES_LIST,
    format_table_cell, add_uit_body_p, add_uit_heading_1, set_table_borders
)
from uit_report_covers import (
    clear_page_border,
    add_cover_page, add_outer_cover_page, add_inner_cover_page
)


def _add_form_line(doc, label_text, is_bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(label_text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.bold = is_bold
    return p


def _add_dotted_lines(doc, count=4):
    for _ in range(count):
        p_dot = doc.add_paragraph()
        p_dot.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_dot.paragraph_format.tab_stops.add_tab_stop(Inches(6.30), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p_dot.paragraph_format.line_spacing = 1.3
        p_dot.paragraph_format.space_before = Pt(2)
        p_dot.paragraph_format.space_after = Pt(2)
        p_dot.paragraph_format.first_line_indent = Inches(0)
        r_dot = p_dot.add_run("\t")
        r_dot.font.name = 'Times New Roman'
        r_dot.font.size = Pt(11)
        r_dot.font.color.rgb = RGBColor(140, 140, 140)


def add_review_forms(doc):
    """Tạo biểu mẫu Nhận xét của Giảng viên hướng dẫn và Phản biện theo quy chế UIT."""
    add_uit_heading_1(doc, "PHIẾU NHẬN XÉT CỦA CÁN BỘ HƯỚNG DẪN")
    _add_form_line(doc, f"Họ và tên Cán bộ hướng dẫn: {UIT_FRONT_MATTER['instructor_name']}")
    _add_form_line(doc, f"Sinh viên thực hiện: {UIT_FRONT_MATTER['student_name']} - MSSV: {UIT_FRONT_MATTER['student_id']} - Lớp: {UIT_FRONT_MATTER['class_name']}")
    _add_form_line(doc, f"Tên đề tài: {UIT_FRONT_MATTER['title_vi']}")

    for label in [
        "1. Tinh thần, thái độ làm việc và tính chủ động của sinh viên:",
        "2. Đánh giá về khối lượng, tiến độ và chất lượng sản phẩm MLOps:",
        "3. Kết quả thử nghiệm trên dữ liệu lớn (442k bản ghi) và tối ưu Optuna:"
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Inches(0)
        r_l = p.add_run(label)
        r_l.font.name = 'Times New Roman'
        r_l.font.size = Pt(11)
        r_l.font.bold = True
        _add_dotted_lines(doc, 4)

    p_score = doc.add_paragraph()
    p_score.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_score.paragraph_format.space_before = Pt(6)
    p_score.paragraph_format.space_after = Pt(2)
    p_score.paragraph_format.first_line_indent = Inches(0)
    p_score.add_run("4. Điểm đánh giá kết luận: ........................................................ / 10.0 điểm").font.bold = True

    p_score_w = doc.add_paragraph()
    p_score_w.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_score_w.paragraph_format.space_before = Pt(1)
    p_score_w.paragraph_format.space_after = Pt(4)
    p_score_w.paragraph_format.first_line_indent = Inches(0)
    r_w = p_score_w.add_run("(Bằng chữ: " + "." * 110 + ")")
    r_w.font.name = 'Times New Roman'
    r_w.font.size = Pt(11)
    r_w.font.italic = True

    p_sig1 = doc.add_paragraph()
    p_sig1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig1.paragraph_format.space_before = Pt(8)
    r_s1 = p_sig1.add_run("TP. Hồ Chí Minh, ngày ...... tháng ...... năm 2026\nCán bộ hướng dẫn\n(Ký và ghi rõ họ tên)\n\n\nThS. Hà Lê Hoài Trung")
    r_s1.font.name = 'Times New Roman'
    r_s1.font.italic = True
    doc.add_page_break()

    add_uit_heading_1(doc, "PHIẾU NHẬN XÉT CỦA CÁN BỘ PHẢN BIỆN")
    p_pb = doc.add_paragraph()
    p_pb.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_pb.paragraph_format.tab_stops.add_tab_stop(Inches(6.30), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p_pb.paragraph_format.line_spacing = 1.3
    p_pb.paragraph_format.space_before = Pt(3)
    p_pb.paragraph_format.space_after = Pt(2)
    p_pb.paragraph_format.first_line_indent = Inches(0)
    r_pb1 = p_pb.add_run("Họ và tên Cán bộ phản biện:")
    r_pb1.font.name = 'Times New Roman'
    r_pb1.font.size = Pt(11)
    r_pb2 = p_pb.add_run("\t")
    r_pb2.font.color.rgb = RGBColor(140, 140, 140)

    _add_form_line(doc, f"Sinh viên thực hiện: {UIT_FRONT_MATTER['student_name']} - MSSV: {UIT_FRONT_MATTER['student_id']}")
    _add_form_line(doc, f"Tên đề tài: {UIT_FRONT_MATTER['title_vi']}")

    for label in [
        "1. Đánh giá tính mới, tính thực tiễn và độ hoàn thiện mã nguồn:",
        "2. Câu hỏi chất vấn phản biện dành cho sinh viên:"
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Inches(0)
        r_l = p.add_run(label)
        r_l.font.name = 'Times New Roman'
        r_l.font.size = Pt(11)
        r_l.font.bold = True
        _add_dotted_lines(doc, 4)

    p_r_score = doc.add_paragraph()
    p_r_score.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_r_score.paragraph_format.space_before = Pt(6)
    p_r_score.paragraph_format.space_after = Pt(2)
    p_r_score.paragraph_format.first_line_indent = Inches(0)
    p_r_score.add_run("3. Điểm đánh giá phản biện: ........................................................ / 10.0 điểm").font.bold = True

    p_r_w = doc.add_paragraph()
    p_r_w.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_r_w.paragraph_format.space_before = Pt(1)
    p_r_w.paragraph_format.space_after = Pt(4)
    p_r_w.paragraph_format.first_line_indent = Inches(0)
    r_rw = p_r_w.add_run("(Bằng chữ: " + "." * 110 + ")")
    r_rw.font.name = 'Times New Roman'
    r_rw.font.size = Pt(11)
    r_rw.font.italic = True

    p_sig2 = doc.add_paragraph()
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig2.paragraph_format.space_before = Pt(8)
    r_s2 = p_sig2.add_run("TP. Hồ Chí Minh, ngày ...... tháng ...... năm 2026\nCán bộ phản biện\n(Ký và ghi rõ họ tên)\n\n\n\n")
    r_s2.font.name = 'Times New Roman'
    r_s2.font.italic = True
    doc.add_page_break()

    add_uit_heading_1(doc, "THÔNG TIN HỘI ĐỒNG BẢO VỆ ĐỒ ÁN")
    add_uit_body_p(doc, f"Hội đồng chấm bảo vệ đồ án môn học, thành lập theo {UIT_FRONT_MATTER['decision_number']}", first_line_indent=False)
    col_widths_b = [Inches(1.8), Inches(2.2), Inches(2.3)]
    t_board = doc.add_table(rows=1, cols=3)
    t_board.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_board, col_widths=col_widths_b)
    format_table_cell(t_board.rows[0].cells[0], "Vai Trò", bold=True, font_size=11)
    format_table_cell(t_board.rows[0].cells[1], "Họ và Tên", bold=True, font_size=11)
    format_table_cell(t_board.rows[0].cells[2], "Đơn Vị Công Tác", bold=True, font_size=11)

    members = [
        ("Chủ tịch Hội đồng", "ThS. Hà Lê Hoài Trung", "Khoa Hệ thống Thông tin - UIT"),
        ("Thư ký Hội đồng", "ThS. Giảng viên Bộ môn", "Khoa Hệ thống Thông tin - UIT"),
        ("Ủy viên Phản biện", "TS. Cán bộ Phản biện", "Khoa Khoa học Máy tính - UIT")
    ]
    for role, name, aff in members:
        row = t_board.add_row()
        format_table_cell(row.cells[0], role, bold=True, font_size=11)
        format_table_cell(row.cells[1], name, bold=False, font_size=11)
        format_table_cell(row.cells[2], aff, bold=False, font_size=11)
    doc.add_page_break()


def add_pledge_and_thanks(doc):
    """Tạo Lời cam đoan và Lời cảm ơn."""
    add_uit_heading_1(doc, "LỜI CAM ĐOAN")
    add_uit_body_p(doc, UIT_FRONT_MATTER["pledge"])
    p_sig_p = doc.add_paragraph()
    p_sig_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sp = p_sig_p.add_run("TP. Hồ Chí Minh, ngày 22 tháng 09 năm 2026\nSinh viên thực hiện\n(Ký và ghi rõ họ tên)\n\n\nNguyễn Trọng Khôi")
    r_sp.font.name = 'Times New Roman'
    r_sp.font.italic = True
    doc.add_page_break()

    add_uit_heading_1(doc, "LỜI CẢM ƠN")
    add_uit_body_p(doc, UIT_FRONT_MATTER["acknowledgments"])
    p_sig_a = doc.add_paragraph()
    p_sig_a.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sa = p_sig_a.add_run("Sinh viên thực hiện\n\nNguyễn Trọng Khôi")
    r_sa.font.name = 'Times New Roman'
    r_sa.font.italic = True
    doc.add_page_break()


def add_abstract_bilingual(doc):
    """Tạo Tóm tắt đề tài (tiếng Việt) và Abstract (tiếng Anh) kèm Từ khóa."""
    add_uit_heading_1(doc, "TÓM TẮT ĐỀ TÀI")
    add_uit_body_p(doc, UIT_FRONT_MATTER["abstract_vi"])
    p_kw_vi = doc.add_paragraph()
    p_kw_vi.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_k1 = p_kw_vi.add_run("Từ khóa: ")
    r_k1.font.bold = True
    r_k1.font.name = 'Times New Roman'
    r_k2 = p_kw_vi.add_run(UIT_FRONT_MATTER["keywords_vi"])
    r_k2.font.italic = True
    r_k2.font.name = 'Times New Roman'
    doc.add_page_break()

    add_uit_heading_1(doc, "ABSTRACT")
    add_uit_body_p(doc, UIT_FRONT_MATTER["abstract_en"])
    p_kw_en = doc.add_paragraph()
    p_kw_en.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_e1 = p_kw_en.add_run("Keywords: ")
    r_e1.font.bold = True
    r_e1.font.name = 'Times New Roman'
    r_e2 = p_kw_en.add_run(UIT_FRONT_MATTER["keywords_en"])
    r_e2.font.italic = True
    r_e2.font.name = 'Times New Roman'
    doc.add_page_break()


def add_toc_and_lists(doc):
    """Tạo Mục lục tổng quan, Danh mục hình ảnh, Danh mục bảng biểu và Danh mục từ viết tắt."""
    add_uit_heading_1(doc, "MỤC LỤC TỔNG QUAN")
    toc_items = [
        ("CHƯƠNG 1: TỔNG QUAN VỀ ĐỀ TÀI", "1"),
        ("  1.1. Phân Loại Bài Toán: Lưu Trữ, Xử Lý & Dự Đoán Khách Hàng Rời Bỏ", "1"),
        ("  1.2. Đặc Tính và Loại Dữ Liệu Dạng Bảng (Structured Tabular Data)", "1"),
        ("  1.3. Quy Mô Dữ Liệu & So Sánh Tốc Độ Xử Lý Máy Chủ Truyền Thống vs Đám Mây", "2"),
        ("  1.4. Phân Loại Các Thành Phần Hệ Thống Theo Mô Hình Dịch Vụ Đám Mây", "2"),
        ("  1.5. Khảo Sát Học Thuật & Đóng Góp Mới Của Đề Tài (2024 - 2026)", "2"),
        ("CHƯƠNG 2: CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ TIÊN TIẾN", "3"),
        ("  2.1. Quy Trình MLOps và Các Cấp Độ Trưởng Thành", "3"),
        ("  2.2. Các Thuật Toán Học Máy Dạng Bảng & Tối Ưu Bayes", "3"),
        ("  2.3. Feast Feature Store & Kiến Trúc Kho Kép Triệt Tiêu Skew", "4"),
        ("  2.4. Lý Thuyết Giám Sát Trôi Dạt Dữ Liệu (Data Drift)", "4"),
        ("  2.5. Khảo Sát Các Mô Hình Dịch Vụ Đám Mây", "4"),
        ("CHƯƠNG 3: PHÂN TÍCH VÀ THIẾT KẾ KIẾN TRÚC HỆ THỐNG", "5"),
        ("  3.1. Phân Tích Yêu Cầu & Rubric Đánh Giá Điểm 10", "5"),
        ("  3.2. Kiến Trúc Phân Tầng ChurnGuard MLOps", "5"),
        ("  3.3. Thiết Kế Luồng Dữ Liệu & Feast Dual-Store", "6"),
        ("  3.4. Thiết Kế CI/CD Pipeline & GitHub Actions Runner", "6"),
        ("CHƯƠNG 4: HIỆN THỰC VÀ THỰC NGHIỆM HỆ THỐNG", "7"),
        ("  4.1. Tiền Xử Lý Dữ Liệu, Khử Trùng Lặp & Data Contracts", "7"),
        ("  4.2. Kết Quả Huấn Luyện Đa Mô Hình & Tối Ưu Siêu Tham Số Optuna", "8"),
        ("  4.3. Phân Tích Lỗi và Trực Quan Hóa Giải Thích SHAP", "9"),
        ("CHƯƠNG 5: ĐÁNH GIÁ ĐỐI CHUẨN ĐÁM MÂY VÀ TỐI ƯU HIỆU NĂNG", "10"),
        ("  5.1. Thử Nghiệm Hiệu Năng Đọc Ghi I/O (Parquet vs CSV vs SQLite)", "10"),
        ("  5.2. Đo Kiểm Độ Trễ Mạng Đa Vùng (Multi-Region Latency)", "10"),
        ("  5.3. Giám Sát Drift Thời Gian Thực & Kích Hoạt Tái Huấn Luyện", "11"),
        ("  5.4. Dự Toán Chi Phí Hạ Tầng Đám Mây (TCO Analysis)", "11"),
        ("CHƯƠNG 6: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", "12"),
        ("  6.1. Các Kết Quả Đạt Được Của Đồ Án", "12"),
        ("  6.2. Phân Tích Giới Hạn Dữ Liệu Tổng Hợp Kaggle", "12"),
        ("  6.3. Hướng Phát Triển Tương Lai", "12"),
        ("TÀI LIỆU THAM KHẢO CHUẨN IEEE", "13"),
        ("PHỤ LỤC", "14")
    ]
    for item, pg in toc_items:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_t.paragraph_format.tab_stops.add_tab_stop(Inches(6.30), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p_t.paragraph_format.line_spacing = 1.15
        p_t.paragraph_format.space_before = Pt(0)
        p_t.paragraph_format.space_after = Pt(1)
        r1 = p_t.add_run(item)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        is_major = (item.startswith("CHƯƠNG") or item.startswith("TÀI LIỆU") or item.startswith("PHỤ LỤC"))
        if is_major:
            r1.font.bold = True
        r2 = p_t.add_run(f"\t{pg}")
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        if is_major:
            r2.font.bold = True
    doc.add_page_break()

    # Danh mục hình ảnh
    add_uit_heading_1(doc, "DANH MỤC CÁC HÌNH ẢNH VÀ BIỂU ĐỒ")
    for fig_id, fig_title, fig_pg in FIGURES_LIST:
        p_f = doc.add_paragraph()
        p_f.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_f.paragraph_format.tab_stops.add_tab_stop(Inches(6.30), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p_f.paragraph_format.line_spacing = 1.2
        p_f.paragraph_format.space_before = Pt(1)
        p_f.paragraph_format.space_after = Pt(1)
        full_title = f"{fig_id}: {fig_title}"
        r_id = p_f.add_run(full_title)
        r_id.font.name = 'Times New Roman'
        r_id.font.size = Pt(11)
        r_p = p_f.add_run(f"\t{fig_pg}")
        r_p.font.name = 'Times New Roman'
        r_p.font.size = Pt(11)
        r_p.font.bold = True
    doc.add_page_break()

    # Danh mục bảng biểu
    add_uit_heading_1(doc, "DANH MỤC CÁC BẢNG")
    for tbl_id, tbl_title, tbl_pg in TABLES_LIST:
        p_tb = doc.add_paragraph()
        p_tb.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_tb.paragraph_format.tab_stops.add_tab_stop(Inches(6.30), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p_tb.paragraph_format.line_spacing = 1.2
        p_tb.paragraph_format.space_before = Pt(1)
        p_tb.paragraph_format.space_after = Pt(1)
        full_title = f"{tbl_id}: {tbl_title}"
        r_id = p_tb.add_run(full_title)
        r_id.font.name = 'Times New Roman'
        r_id.font.size = Pt(11)
        r_p = p_tb.add_run(f"\t{tbl_pg}")
        r_p.font.name = 'Times New Roman'
        r_p.font.size = Pt(11)
        r_p.font.bold = True
    doc.add_page_break()

    # Danh mục từ viết tắt
    add_uit_heading_1(doc, "DANH MỤC CÁC TỪ VIẾT TẮT")
    col_widths_abbr = [Inches(0.6), Inches(1.4), Inches(2.3), Inches(2.0)]
    t_abbr = doc.add_table(rows=1, cols=4)
    t_abbr.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_abbr, col_widths=col_widths_abbr)
    format_table_cell(t_abbr.rows[0].cells[0], "STT", bold=True, font_size=10)
    format_table_cell(t_abbr.rows[0].cells[1], "Từ Viết Tắt", bold=True, font_size=10)
    format_table_cell(t_abbr.rows[0].cells[2], "Thuật Ngữ Tiếng Anh", bold=True, font_size=10)
    format_table_cell(t_abbr.rows[0].cells[3], "Giải Thích Ý Nghĩa", bold=True, font_size=10)
    for stt, abbr, en, vi in ABBREVIATIONS_LIST:
        row = t_abbr.add_row()
        format_table_cell(row.cells[0], stt, bold=False, font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_table_cell(row.cells[1], abbr, bold=True, font_size=10)
        format_table_cell(row.cells[2], en, bold=False, font_size=10, italic=True)
        format_table_cell(row.cells[3], vi, bold=False, font_size=10)
    # Không gọi doc.add_page_break() ở đây vì generate_docx_report thêm sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
